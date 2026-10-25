"""The free spreadsheet that new subscribers get by email: site/free-list/.

Built from the same free-only rows as the website (build_site.py's `public`), so
member columns can't reach it. Two files per sale, CSV and Excel, plus stable
links that always point at the current sale:
  /free-list/latest.xlsx, /free-list/latest.csv   (site/_redirects)
The folder isn't linked from any page and is marked noindex (site/_headers): the
link is meant to travel in the beehiiv welcome email and the weekly issue.
"""
import csv
import glob
import os
from datetime import date
from urllib.parse import quote_plus

SITE_URL = "https://firsttuesdaylist.com"
TYPES = {"single_family": "Single family", "condo_townhome": "Condo / townhome", "multifamily": "Multifamily",
         "commercial": "Commercial", "vacant_land": "Vacant land"}
HEADER = ["Notice ID", "Address", "City", "ZIP", "Property type", "Sale date", "Sale time window", "Sale location",
          "Notice type", "FEMA flood zone", "High-risk flood area (SFHA)", "Posted more than once", "First seen",
          "Property page", "Notice PDF", "Google Maps"]
LINK_COLS = {"Property page", "Notice PDF", "Google Maps"}


def rows_for(sale, public):
    out = []
    for r in sorted(public, key=lambda r: (r["zip"], r["addr"])):
        maps = (f"https://www.google.com/maps/search/?api=1&query={r['lat']},{r['lon']}" if r["lat"] is not None
                else "https://www.google.com/maps/search/?api=1&query=" + quote_plus(f"{r['addr']}, {r['city']}, TX {r['zip']}"))
        out.append([r["id"], r["addr"], r["city"], r["zip"], TYPES.get(r["type"], "Other"), sale, r["window"] or "",
                    r["location"] or "", "HOA lien" if r["notice"] == "hoa" else "Mortgage foreclosure", r["zone"] or "",
                    "Yes" if r["sfha"] else "No", "Yes" if r.get("repeat") else "No", r["first"],
                    f"{SITE_URL}/property/{r['id']}", f"{SITE_URL}/notice/{r['id']}" if r["notice_url"] else "", maps])
    return out


def write_free_list(site, sale, public, label_year):
    out = os.path.join(site, "free-list")
    os.makedirs(out, exist_ok=True)
    for old in glob.glob(os.path.join(out, "first-tuesday-list-*")):  # only the current sale is offered
        os.remove(old)
    base = f"first-tuesday-list-{sale}"
    rows = rows_for(sale, public)

    with open(os.path.join(out, base + ".csv"), "w", newline="", encoding="utf-8-sig") as f:  # BOM: Excel opens it as UTF-8
        w = csv.writer(f)
        w.writerow(HEADER)
        w.writerows(rows)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = f"{label_year} sale"[:31]
    ws.append(HEADER)
    for row in rows:
        ws.append(row)
    navy, white = PatternFill("solid", fgColor="13294B"), Font(bold=True, color="FFFFFF")
    for c in ws[1]:
        c.fill, c.font, c.alignment = navy, white, Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    widths = {"Notice ID": 16, "Address": 30, "City": 14, "ZIP": 8, "Property type": 17, "Sale date": 11, "Sale time window": 18,
              "Sale location": 40, "Notice type": 21, "FEMA flood zone": 10, "High-risk flood area (SFHA)": 12,
              "Posted more than once": 11, "First seen": 11, "Property page": 13, "Notice PDF": 12, "Google Maps": 12}
    for i, h in enumerate(HEADER, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths[h]
        if h in LINK_COLS:
            for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for c in cell:
                    if c.value:
                        c.hyperlink, c.value, c.style = c.value, {"Property page": "Photos and map", "Notice PDF": "Save PDF",
                                                                  "Google Maps": "Open map"}[h], "Hyperlink"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    about = wb.create_sheet("Read me")
    for line in (
        [f"First Tuesday List · Harris County foreclosure auction · {label_year}"],
        [f"{len(rows)} properties. Built {date.today().isoformat()} from notices filed with the Harris County Clerk."],
        [""],
        ["Every property has a page with aerial photos, Street View, the map and the county notice:"],
        [f"{SITE_URL}/map"],
        [""],
        ["Members also get each property's value, loan, estimated balance and equity, and the lender. See " + SITE_URL + "/pricing"],
        [""],
        ["Compiled from public records. Not verified. Do your own title and lien research before bidding."],
        ["First Tuesday List is not a law firm, title company, broker or auctioneer and does not give legal or investment advice."],
    ):
        about.append(line)
    about["A1"].font = Font(bold=True, size=14)
    about.column_dimensions["A"].width = 110
    wb.save(os.path.join(out, base + ".xlsx"))

    # stable links for the emails, and keep the folder out of search engines
    with open(os.path.join(site, "_redirects"), "w", encoding="utf-8") as f:
        f.write(f"/free-list/latest.xlsx /free-list/{base}.xlsx 302\n/free-list/latest.csv /free-list/{base}.csv 302\n")
    with open(os.path.join(site, "_headers"), "w", encoding="utf-8") as f:
        f.write("/free-list/*\n  X-Robots-Tag: noindex, nofollow\n  Cache-Control: no-cache\n")
    return len(rows)
